"""Parse the TOP 1C price PDF into product lines (shared image groups)."""

from __future__ import annotations

import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

import pymupdf

CAT_RE = re.compile(r"^(\d{2}\.\d{2}(?:\.\d)?)\s+(.+)$")
PRICE_RE = re.compile(r"^(\d{1,3}(?: \d{3})*,\d{2})$")
SKU_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9\-/\.]{1,24}$")
# 1C often types Latin SKU suffixes with Cyrillic lookalikes (А/В/С…).
_SKU_CYR_TO_LAT = str.maketrans(
    "АВЕКМНОРСТХавекмнорстх",
    "ABEKMHOPCTXabekmhopctx",
)


def latinize_sku(text: str) -> str:
    return text.translate(_SKU_CYR_TO_LAT).strip()


# Card geometry — keep in sync with catalog.render (2×6 grid)
_CELL_H = (790.5 - 52.0 - 5 * 5.6) / 6
_ROW_GAP = 5.6
_ROW_H = 11.0
_HEAD_H = 7.5
_TITLE_PAD = 26.0
_PAD_BOTTOM = 4.0
_MIN_STACK_IMG = 56.0
_SLOTS_PER_COL = 6


def max_split_rows() -> int:
    """How many variant rows fit beside the photo in a single 2×6 cell."""
    avail = _CELL_H - _TITLE_PAD - _PAD_BOTTOM
    return max(1, int((avail - _HEAD_H) / _ROW_H))


def stacked_slot_count(n: int) -> int:
    """Slots for a card: split single while the table fits, else stacked 2-col."""
    if n <= max_split_rows():
        return 1
    rows = (n + 1) // 2
    table_h = _HEAD_H + rows * _ROW_H
    need = _TITLE_PAD + _MIN_STACK_IMG + 4.0 + table_h
    for k in range(2, _SLOTS_PER_COL + 1):
        if k * _CELL_H + (k - 1) * _ROW_GAP >= need:
            return k
    return _SLOTS_PER_COL


def uses_two_col_table(n: int) -> bool:
    """Grown (2+ slot) cards put variants in two columns."""
    return stacked_slot_count(n) > 1


def product_line_slots(n: int) -> int:
    return stacked_slot_count(n)


def clean_name(name: str, sku: str = "") -> str:
    name = name.replace("*СЛЕДОПЫТ", "«СЛЕДОПЫТ").replace('СЛЕДОПЫТ"', "СЛЕДОПЫТ»")
    name = re.sub(r'"\s*Рос+омаха\s*"', "«РОСОМАХА»", name, flags=re.I)
    name = re.sub(r'"\s*БУЛЬДОЗЕР\s*"', "«БУЛЬДОЗЕР»", name, flags=re.I)
    name = name.replace("натурайльного", "натурального")
    name = name.replace("керамогниту", "керамограниту")
    name = name.replace("камню. кирпичу", "камню, кирпичу")
    name = re.sub(r"SDS\s*(?:-?\s*plus|\+)", "SDS+", name, flags=re.I)
    name = re.sub(r"(?<=[а-яёА-ЯЁ])\.(?=[а-яёА-ЯЁ])", ". ", name)
    name = re.sub(r"(?<!\s)«", " «", name)
    name = re.sub(r"\s+", " ", name).strip()
    name = re.sub(r"^\d{4,5}\s+", "", name)
    if sku:
        name = re.sub(rf"\s*{re.escape(sku)}\b", "", name)
    name = re.sub(r"(?:«РОСОМАХА»\s*)+", "«РОСОМАХА» ", name)
    name = re.sub(r"\s*TOLSEN\s+\d{4,5}\s*$", " TOLSEN", name)
    name = re.sub(r"ПРАКТИКА\"", "ПРАКТИКА «", name)
    name = re.sub(r"ммTOLSEN", "мм TOLSEN", name)
    # space after comma, but keep decimal commas: 0,5л / 22,23 мм
    name = re.sub(r",(?!\d)(?=\S)", ", ", name)
    name = re.sub(r"\s+", " ", name).strip(" ,.")
    return name


@dataclass
class Product:
    name: str
    sku: str
    price: str
    category_code: str
    category_name: str
    image_xref: int | None = None
    image_page: int | None = None


@dataclass
class ProductLine:
    category_code: str
    category_name: str
    products: list[Product] = field(default_factory=list)
    image_xref: int | None = None
    image_page: int | None = None
    image_png: bytes | None = None
    excel_row: int | None = None
    excel_source: str | None = None
    excel_image: bytes | None = None

    @property
    def slots(self) -> int:
        return product_line_slots(len(self.products))


def _cluster_rows(items: list[dict], y_tol: float = 4.0) -> list[list[dict]]:
    items = sorted(items, key=lambda i: (i["y"], i["x"]))
    rows: list[list[dict]] = []
    for it in items:
        if rows and abs(it["y"] - rows[-1][0]["y"]) <= y_tol:
            rows[-1].append(it)
        else:
            rows.append([it])
    return rows


def parse_price(pdf_path: str | Path) -> list[ProductLine]:
    doc = pymupdf.open(pdf_path)
    current_cat = ("", "")
    lines: list[ProductLine] = []
    current: ProductLine | None = None
    pending_name: list[str] = []

    def flush_pending():
        nonlocal pending_name
        pending_name = []

    def start_line(page_i: int, xref: int | None):
        nonlocal current
        current = ProductLine(
            category_code=current_cat[0],
            category_name=current_cat[1],
            image_xref=xref,
            image_page=page_i,
        )
        lines.append(current)

    for page_i, page in enumerate(doc):
        spans: list[dict] = []
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                text = "".join(s["text"] for s in line["spans"]).strip()
                if not text:
                    continue
                x0, y0, x1, y1 = line["bbox"]
                spans.append({"text": text, "x": x0, "x1": x1, "y": y0})

        images = []
        for info in page.get_images(full=True):
            xref = info[0]
            for rect in page.get_image_rects(xref):
                images.append({"xref": xref, "y0": rect.y0, "y1": rect.y1, "x0": rect.x0})
        images.sort(key=lambda im: im["y0"])

        def image_at(y: float) -> int | None:
            # 1C groups align the photo top with the 12345 marker.
            # Matching by overlap with slack stole the previous group's image.
            if not images:
                return None
            best = min(images, key=lambda im: abs(im["y0"] - y))
            if abs(best["y0"] - y) > 48:
                return None
            return best["xref"]

        rows = _cluster_rows(spans)
        for row in rows:
            texts = [t["text"] for t in sorted(row, key=lambda t: t["x"])]
            joined = " ".join(texts)
            cat_m = CAT_RE.match(joined) or (
                CAT_RE.match(texts[0]) if len(texts) == 1 else None
            )
            if cat_m:
                current_cat = (cat_m.group(1), cat_m.group(2).strip())
                current = None
                flush_pending()
                continue

            has_marker = any(t["text"] == "12345" for t in row)
            price_items = [t for t in row if PRICE_RE.match(t["text"]) and t["x"] > 500]
            sku_items = [
                t
                for t in row
                if t["x"] > 430
                and t["x"] < 530
                and SKU_RE.match(latinize_sku(t["text"]))
                and t["text"] != "12345"
            ]
            name_items = [
                t
                for t in row
                if 150 < t["x"] < 430 and t["text"] not in {"12345"}
            ]

            # SKU glued to the end of a long name
            if not sku_items and name_items:
                m = re.search(
                    r"\s([A-Za-z0-9АВЕКМНОРСТХавекмнорстх][A-Za-z0-9АВЕКМНОРСТХавекмнорстх\-/\.]{2,20})$",
                    name_items[0]["text"],
                )
                if m and price_items:
                    sku_items = [{"text": m.group(1), "x": 488, "y": name_items[0]["y"]}]
                    name_items[0]["text"] = name_items[0]["text"][: m.start()].strip()

            if has_marker:
                xref = image_at(row[0]["y"])
                start_line(page_i, xref)
                flush_pending()

            if name_items and not price_items and not sku_items:
                extra = " ".join(t["text"] for t in name_items).strip()
                if (
                    current
                    and current.products
                    and re.search(
                        r"^([\"«]?РОСОМАХА|проволоки|уп\)|шт\)|,?\s*\d+\s*шт)",
                        extra,
                        re.I,
                    )
                ):
                    current.products[-1].name = clean_name(
                        current.products[-1].name + " " + extra,
                        current.products[-1].sku,
                    )
                else:
                    pending_name.append(extra)
                continue

            if price_items and (sku_items or name_items or pending_name):
                name_parts = pending_name + [t["text"] for t in name_items]
                sku = latinize_sku(sku_items[0]["text"]) if sku_items else ""
                price = price_items[0]["text"]
                name = clean_name(" ".join(name_parts), sku)
                if current is None:
                    xref = image_at(row[0]["y"])
                    start_line(page_i, xref)
                if sku:
                    current.products.append(
                        Product(
                            name=name,
                            sku=sku,
                            price=price,
                            category_code=current.category_code,
                            category_name=current.category_name,
                            image_xref=current.image_xref,
                            image_page=current.image_page,
                        )
                    )
                flush_pending()
            elif name_items:
                pending_name.extend(t["text"] for t in name_items)

    doc.close()
    for ln in lines:
        if any("«РОСОМАХА»" in p.name for p in ln.products):
            for p in ln.products:
                if "«РОСОМАХА»" not in p.name:
                    p.name = (p.name + " «РОСОМАХА»").strip()
    return [ln for ln in lines if ln.products]


def format_excel_price(value) -> str:
    if isinstance(value, (int, float)):
        s = f"{float(value):.2f}".replace(".", ",")
        whole, frac = s.split(",")
        if len(whole) > 3:
            parts: list[str] = []
            while whole:
                parts.append(whole[-3:])
                whole = whole[:-3]
            whole = " ".join(reversed(parts))
        return f"{whole},{frac}"
    text = str(value).strip()
    return text


def _excel_images_by_row(xlsx_path: str | Path) -> dict[int, bytes]:
    """Map 1-based Excel row of a 12345 photo cell to PNG bytes."""
    z = zipfile.ZipFile(xlsx_path)
    rel_xml = z.read("xl/drawings/_rels/drawing1.xml.rels").decode()
    rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rel_xml))
    draw = z.read("xl/drawings/drawing1.xml")
    if draw.startswith(b"\xef\xbb\xbf"):
        draw = draw[3:]
    root = ET.fromstring(draw)

    def local(tag: str) -> str:
        return tag.split("}")[-1] if "}" in tag else tag

    out: dict[int, bytes] = {}
    for child in list(root):
        if local(child.tag) != "twoCellAnchor":
            continue
        frm = None
        embed = None
        for el in child.iter():
            if local(el.tag) == "from":
                for sub in el:
                    if local(sub.tag) == "row" and frm is None:
                        frm = int(sub.text)
            if local(el.tag) == "blip":
                embed = el.attrib.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                )
        if frm is None or not embed:
            continue
        target = rels.get(embed)
        if not target:
            continue
        data = z.read("xl/media/" + Path(target).name)
        out[frm + 1] = data
    return out


def _excel_columns(ws) -> tuple[int, int, int]:
    """Name / SKU / price columns. New 1C export: M/O/P; older file: F/N/O."""
    name_c = sku_c = price_c = None
    for r in range(1, 10):
        for c in range(1, 17):
            lab = str(ws.cell(r, c).value or "").strip().lower()
            if lab in {"номенклатура", "наименование"}:
                name_c = c
            elif lab in {"код", "артикул"}:
                sku_c = c
            elif lab == "цена":
                price_c = c
    if name_c and sku_c and price_c:
        return name_c, sku_c, price_c
    return 6, 14, 15


def parse_excel(xlsx_path: str | Path) -> list[ProductLine]:
    """Read the 1C Excel price: one product line per 12345 photo group."""
    from openpyxl import load_workbook

    xlsx_path = Path(xlsx_path)
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    name_c, sku_c, price_c = _excel_columns(ws)
    images = _excel_images_by_row(xlsx_path)
    skip_sku = {"код", "kod", "sku", "артикул"}
    skip_name = {"номенклатура", "наименование"}
    lines: list[ProductLine] = []
    current: ProductLine | None = None
    current_cat = ("", "")
    source = xlsx_path.name
    for r in range(1, (ws.max_row or 0) + 1):
        a = str(ws.cell(r, 1).value or "").strip()
        name_raw = str(ws.cell(r, name_c).value or "").strip()
        sku_raw = str(ws.cell(r, sku_c).value or "").strip()
        sku = latinize_sku(sku_raw)
        price_val = ws.cell(r, price_c).value
        cat_m = CAT_RE.match(a)
        if cat_m:
            current_cat = (cat_m.group(1), cat_m.group(2).strip())
            current = None
            continue
        if a.replace(" ", "") == "12345":
            current = ProductLine(
                category_code=current_cat[0],
                category_name=current_cat[1],
                excel_row=r,
                excel_source=source,
                excel_image=images.get(r),
            )
            lines.append(current)
        if not name_raw or not sku:
            continue
        if sku.lower() in skip_sku or name_raw.lower() in skip_name:
            continue
        if current is None:
            current = ProductLine(
                category_code=current_cat[0],
                category_name=current_cat[1],
                excel_row=r,
                excel_source=source,
                excel_image=images.get(r),
            )
            lines.append(current)
        name = clean_name(name_raw, sku)
        current.products.append(
            Product(
                name=name,
                sku=sku,
                price=format_excel_price(price_val),
                category_code=current.category_code,
                category_name=current.category_name,
            )
        )
    for ln in lines:
        if any("«РОСОМАХА»" in p.name for p in ln.products):
            for p in ln.products:
                if "«РОСОМАХА»" not in p.name:
                    p.name = (p.name + " «РОСОМАХА»").strip()
    return [ln for ln in lines if ln.products]


def parse_excels(paths: list[str | Path]) -> list[ProductLine]:
    lines: list[ProductLine] = []
    for path in paths:
        lines.extend(parse_excel(path))
    return lines


def attach_pdf_images(lines: list[ProductLine], pdf_path: str | Path) -> None:
    """Keep catalog photos from the 1C PDF, matched by SKU to Excel groups."""
    pdf_lines = parse_price(pdf_path)
    sku_img: dict[str, tuple[int | None, int | None]] = {}
    for ln in pdf_lines:
        key = (ln.image_page, ln.image_xref)
        for p in ln.products:
            sku_img[p.sku] = key
    for ln in lines:
        hits = [sku_img[p.sku] for p in ln.products if p.sku in sku_img]
        if not hits:
            continue
        page, xref = Counter(hits).most_common(1)[0][0]
        ln.image_page = page
        ln.image_xref = xref
        for p in ln.products:
            p.image_page = page
            p.image_xref = xref


# Groups that must use the Excel photo: previous metal-disc override plus
# the three cards that visually changed in 1.xlsx.
EXCEL_PHOTO_SKUS = {
    "408125",  # отрезные по металлу и нерж. стали «РОСОМАХА»
    "100125",  # алмазный универсальный «РОСОМАХА» (125 мм вместо 230)
    "33269A",  # зубило SDS+
    "031-211",  # лента шлифовальная ПРАКТИКА 100×610
}


def apply_excel_photo_overrides(lines: list[ProductLine], xlsx_path: str | Path | None = None) -> None:
    """Excel photos for new cards (no PDF) and for the few visual replacements."""
    extra: dict[int, bytes] = {}
    if xlsx_path:
        extra = _excel_images_by_row(xlsx_path)
    for ln in lines:
        blob = ln.excel_image or extra.get(ln.excel_row or -1)
        if not blob:
            continue
        skus = {p.sku for p in ln.products}
        if ln.image_xref is None or skus & EXCEL_PHOTO_SKUS:
            ln.image_png = blob


if __name__ == "__main__":
    src = Path("/workspace/Прайс ТОП до ЛЕЗВИЙ.pdf")
    parsed = parse_price(src)
    print(f"lines={len(parsed)} products={sum(len(l.products) for l in parsed)}")
    for ln in parsed:
        if ln.category_code.startswith("01") or ln.category_code.startswith("02"):
            print(
                f"{ln.category_code} slots={ln.slots} n={len(ln.products)} "
                f"img={ln.image_xref} | {ln.products[0].name[:70]}"
            )
            for p in ln.products:
                print(f"    {p.sku:12} {p.price:>10}  {p.name[:90]}")
